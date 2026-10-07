from pathlib import Path
from datetime import datetime
from typing import List
import io

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_user, require_role
from app.core.constants import Roles
from app.models import (
    Usuario,
    ProcesoEtapaProductiva,
    Ficha,
    NotificacionMensaje,
    EstadoEnvioEmail,
)
from app.modules.cartas_presentacion import crud, schemas

# openpyxl para generar Excel
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False


router = APIRouter(prefix="/cartas-presentacion", tags=["Cartas de Presentación"])

CARTAS_DIR = Path("uploads/cartas")
CARTAS_DIR.mkdir(parents=True, exist_ok=True)

MAX_PDF_SIZE = 10 * 1024 * 1024  # 10 MB


# ============================================================
# HELPER: Crear notificación para el aprendiz
# ============================================================

def _crear_notificacion_carta(
    db: Session,
    aprendiz_id: int,
    coordinador: Usuario,
):
    """Crea una notificación al aprendiz cuando se sube una carta."""
    try:
        nombre_coord = f"{coordinador.nombre} {coordinador.apellido}".strip()

        notif = NotificacionMensaje(
            remitente_usuario_id=coordinador.id,
            destinatario_usuario_id=aprendiz_id,
            asunto="Nueva carta de presentación disponible",
            cuerpo=(
                f"{nombre_coord} subió tu carta de presentación. "
                f"Ya puedes verla y descargarla desde 'Mis Cartas de Presentación'."
            ),
            estado_envio_email=EstadoEnvioEmail.PENDIENTE,
            is_active=True,
        )
        db.add(notif)
        db.commit()
        db.refresh(notif)
        print(f"✅ Notificación creada para aprendiz {aprendiz_id}")
        return notif
    except Exception as e:
        print(f"[CARTAS] Error al crear notificación: {e}")
        return None


# ============================================================
# LISTAR
# ============================================================

@router.get("/", response_model=schemas.CartasListOut)
def listar(
    proceso_id: int = None,
    aprendiz_id: int = None,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
):
    """
    - Coordinador/Admin: ve todas las cartas
    - Aprendiz: solo las suyas
    """
    if current_user.rol_id == Roles.APRENDIZ:
        cartas = crud.listar_cartas(db, aprendiz_id=current_user.id)
    else:
        cartas = crud.listar_cartas(db, proceso_id=proceso_id, aprendiz_id=aprendiz_id)

    return {
        "cartas": [crud.serializar_carta(db, c) for c in cartas],
        "total": len(cartas),
    }


# ============================================================
# PLANTILLA EXCEL DE APRENDICES
# ============================================================

@router.get("/plantilla-aprendices")
def descargar_plantilla_aprendices(
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role(Roles.ADMIN, Roles.COORDINADOR)),
):
    """Genera un Excel con los aprendices activos y el nombre sugerido para sus PDFs."""
    if not OPENPYXL_DISPONIBLE:
        raise HTTPException(
            status_code=500,
            detail="openpyxl no está instalado. Ejecuta: pip install openpyxl",
        )

    procesos = db.query(ProcesoEtapaProductiva).filter(
        ProcesoEtapaProductiva.is_active == True,
        ProcesoEtapaProductiva.estado == "ACTIVO",
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Aprendices"

    # Encabezados
    headers = ["#", "Ficha", "Documento", "Nombre del aprendiz", "Nombre sugerido del PDF"]
    ws.append(headers)

    # Estilos del encabezado
    header_fill = PatternFill(start_color="3CA203", end_color="3CA203", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Datos
    contador = 1
    for p in procesos:
        aprendiz = db.query(Usuario).filter(Usuario.id == p.aprendiz_id).first()
        if not aprendiz:
            continue

        ficha = p.ficha.numero_ficha if p.ficha else "SIN_FICHA"
        documento = aprendiz.documento_identidad or "SIN_DOCUMENTO"
        nombre = f"{aprendiz.nombre} {aprendiz.apellido}".strip()
        nombre_pdf = f"{ficha}_{documento}.pdf"

        ws.append([contador, ficha, documento, nombre, nombre_pdf])
        contador += 1

    # Ajustar ancho de columnas
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 40

    # Fila adicional con instrucciones
    ws.append([])
    ws.append(["INSTRUCCIONES:"])
    ws.append(["1. Renombra cada PDF usando la columna 'Nombre sugerido del PDF'"])
    ws.append(["2. Sube todos los PDFs juntos desde el botón 'Carga Masiva'"])
    ws.append(["3. El sistema asignará automáticamente cada PDF a su aprendiz"])

    # Guardar en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_cartas_aprendices.xlsx"},
    )


# ============================================================
# SUBIR PDF INDIVIDUAL (con notificación automática)
# ============================================================

@router.post("/subir", status_code=status.HTTP_201_CREATED)
def subir_carta(
    proceso_id: int,
    file: UploadFile = File(..., description="PDF firmado"),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role(Roles.ADMIN, Roles.COORDINADOR)),
):
    """Coordinador sube el PDF firmado + crea notificación al aprendiz."""
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Solo se aceptan archivos PDF",
        )

    content = file.file.read()
    if len(content) > MAX_PDF_SIZE:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El PDF no puede superar los 10 MB",
        )

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre_seguro = f"carta_proceso_{proceso_id}_{timestamp}.pdf"
    filepath = CARTAS_DIR / nombre_seguro

    with open(filepath, "wb") as buffer:
        buffer.write(content)

    archivo_url = f"/uploads/cartas/{nombre_seguro}"

    carta = crud.crear_carta(
        db=db,
        proceso_id=proceso_id,
        coordinador_id=current_user.id,
        archivo_url=archivo_url,
        nombre_archivo=file.filename,
        tamano_bytes=len(content),
    )

    # 🔔 Notificación al aprendiz
    proceso = db.query(ProcesoEtapaProductiva).filter(
        ProcesoEtapaProductiva.id == proceso_id
    ).first()

    if proceso:
        _crear_notificacion_carta(
            db,
            aprendiz_id=proceso.aprendiz_id,
            coordinador=current_user,
        )

    return {
        "message": "Carta de presentación subida exitosamente",
        "carta": crud.serializar_carta(db, carta),
    }


# ============================================================
# SUBIR MASIVO (varios PDFs)
# ============================================================

@router.post("/subir-masivo", status_code=status.HTTP_201_CREATED)
def subir_cartas_masivo(
    files: List[UploadFile] = File(..., description="Lista de PDFs"),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role(Roles.ADMIN, Roles.COORDINADOR)),
):
    """
    Sube varios PDFs a la vez.
    Formato del nombre: {FICHA}_{DOCUMENTO}.pdf
    Ejemplo: 2875901_1234567890.pdf
    """
    exitosos = []
    errores = []

    for file in files:
        try:
            # Validar extensión
            if not file.filename.lower().endswith(".pdf"):
                errores.append({
                    "archivo": file.filename,
                    "error": "No es un archivo PDF",
                })
                continue

            # Parsear nombre: FICHA_DOCUMENTO.pdf
            nombre_base = file.filename.rsplit(".", 1)[0]
            partes = nombre_base.split("_")

            if len(partes) < 2:
                errores.append({
                    "archivo": file.filename,
                    "error": "Formato incorrecto. Debe ser FICHA_DOCUMENTO.pdf",
                })
                continue

            ficha_numero = partes[0].strip()
            documento = partes[1].strip()

            # Buscar ficha
            ficha = db.query(Ficha).filter(
                Ficha.numero_ficha == ficha_numero,
                Ficha.is_active == True,
            ).first()

            if not ficha:
                errores.append({
                    "archivo": file.filename,
                    "error": f"No existe la ficha {ficha_numero}",
                })
                continue

            # Buscar aprendiz por documento
            aprendiz = db.query(Usuario).filter(
                Usuario.documento_identidad == documento,
                Usuario.is_active == True,
            ).first()

            if not aprendiz:
                errores.append({
                    "archivo": file.filename,
                    "error": f"No existe aprendiz con documento {documento}",
                })
                continue

            # Buscar proceso activo
            proceso = db.query(ProcesoEtapaProductiva).filter(
                ProcesoEtapaProductiva.aprendiz_id == aprendiz.id,
                ProcesoEtapaProductiva.ficha_id == ficha.id,
                ProcesoEtapaProductiva.is_active == True,
            ).first()

            if not proceso:
                errores.append({
                    "archivo": file.filename,
                    "error": f"{aprendiz.nombre} {aprendiz.apellido} no tiene proceso activo en la ficha {ficha_numero}",
                })
                continue

            # Leer y validar tamaño
            content = file.file.read()
            if len(content) > MAX_PDF_SIZE:
                errores.append({
                    "archivo": file.filename,
                    "error": "El PDF supera los 10 MB",
                })
                continue

            # Guardar archivo
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            nombre_seguro = f"carta_{ficha_numero}_{documento}_{timestamp}.pdf"
            filepath = CARTAS_DIR / nombre_seguro

            with open(filepath, "wb") as buffer:
                buffer.write(content)

            archivo_url = f"/uploads/cartas/{nombre_seguro}"

            # Crear carta
            carta = crud.crear_carta(
                db=db,
                proceso_id=proceso.id,
                coordinador_id=current_user.id,
                archivo_url=archivo_url,
                nombre_archivo=file.filename,
                tamano_bytes=len(content),
            )

            # Notificación al aprendiz
            _crear_notificacion_carta(
                db,
                aprendiz_id=aprendiz.id,
                coordinador=current_user,
            )

            exitosos.append({
                "archivo": file.filename,
                "aprendiz": f"{aprendiz.nombre} {aprendiz.apellido}",
                "ficha": ficha_numero,
                "carta_id": carta.id,
            })

        except Exception as e:
            errores.append({
                "archivo": file.filename,
                "error": f"Error inesperado: {str(e)}",
            })

    return {
        "total": len(files),
        "exitosos": len(exitosos),
        "errores": len(errores),
        "detalles_exitosos": exitosos,
        "detalles_errores": errores,
    }


# ============================================================
# ACTUALIZAR ESTADO (aprobar/rechazar)
# ============================================================

@router.patch("/{carta_id}/estado")
def actualizar_estado(
    carta_id: int,
    payload: schemas.ActualizarEstadoCarta,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role(Roles.ADMIN, Roles.COORDINADOR)),
):
    carta = crud.get_carta(db, carta_id)
    if not carta:
        raise HTTPException(status_code=404, detail="Carta no encontrada")

    if payload.estado not in ("APROBADA", "RECHAZADA", "CARGADA"):
        raise HTTPException(status_code=400, detail="Estado no válido")

    carta = crud.actualizar_estado(db, carta, payload.estado, payload.observaciones)
    return {"message": "Estado actualizado", "carta": crud.serializar_carta(db, carta)}


# ============================================================
# ELIMINAR
# ============================================================

@router.delete("/{carta_id}")
def eliminar(
    carta_id: int,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role(Roles.ADMIN, Roles.COORDINADOR)),
):
    carta = crud.get_carta(db, carta_id)
    if not carta:
        raise HTTPException(status_code=404, detail="Carta no encontrada")

    crud.eliminar_carta(db, carta)
    return {"message": "Carta eliminada"}